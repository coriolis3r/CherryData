# main.py
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
import pandas as pd
import matplotlib

# Obliga a matplotlib a no usar interfaces gráficas de escritorio (crucial para servidores en la nube)
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
import io

# Importaciones de tus módulos locales
import config
import report_utils

app = FastAPI(title="Procesador Avanzado de Calidad de Energía")

# Habilitar CORS para conectar directamente con tu app de React
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.post("/analizar-historico")
@app.post("/analizar-historico/")
async def analizar_historico(archivo: UploadFile = File(...)):
    try:
        # 1. CARGAR CSV DESDE MEMORIA RAM
        contenido = await archivo.read()
        df = pd.read_csv(io.BytesIO(contenido))

        # Verificar columnas válidas según tu config.py
        columnas_validas = [c for c in config.columnas_seleccionadas if c in df.columns]
        if not columnas_validas:
            raise HTTPException(status_code=400,
                                detail="No se encontraron las columnas seleccionadas en el archivo CSV.")

        # 2. PROCESAMIENTO DE DATOS
        df[columnas_validas] = df[columnas_validas].apply(pd.to_numeric, errors="coerce")

        if "date" not in df.columns or "time" not in df.columns:
            raise HTTPException(status_code=400, detail="Faltan las columnas 'date' o 'time' en el CSV.")

        datetime_col = pd.to_datetime(
            df["date"].astype(str) + " " + df["time"].astype(str),
            errors="coerce"
        )
        df = pd.concat([df, datetime_col.rename("datetime")], axis=1)
        df = df.dropna(subset=["datetime"]).sort_values("datetime")

        # 3. TABLA GENERAL DE ESTADÍSTICAS (Pandas DataFrame)
        estadisticas = pd.DataFrame({
            "Mínimo": df[columnas_validas].min(),
            "Máximo": df[columnas_validas].max(),
            "Promedio": df[columnas_validas].mean(),
            "Desv. estándar": df[columnas_validas].std(),
            "P95": df[columnas_validas].quantile(0.95),
            "P99": df[columnas_validas].quantile(0.99)
        }).round(2)

        # 4. GENERACIÓN DE GRÁFICAS EN MEMORIA RAM
        # Guardaremos los flujos binarios de los gráficos aquí
        imagenes_memoria = {}

        for bloque in config.configuracion_reporte:
            plt.figure(figsize=(12, 5))
            grafico_algo = False

            for col_csv, etiqueta in bloque["variables"]:
                if col_csv in df.columns:
                    plt.plot(df["datetime"], df[col_csv], label=etiqueta)
                    grafico_algo = True

            if grafico_algo:
                plt.title(bloque["titulo"])
                plt.xlabel("Fecha y hora")
                plt.ylabel(bloque["ylabel"])
                plt.legend()
                plt.grid(True)
                plt.xticks(rotation=45)
                plt.tight_layout()

                # OPTIMIZACIÓN EN MEMORIA: Guardar la gráfica en un buffer binario en vez del disco duro
                img_buf = io.BytesIO()
                plt.savefig(img_buf, format='png', dpi=250, bbox_inches="tight")
                img_buf.seek(0)

                # Asociamos el buffer con el identificador del bloque
                imagenes_memoria[bloque["titulo"]] = img_buf

            plt.close()

        # 5. CÁLCULO DINÁMICO DE LOS INDICADORES DEL RESUMEN EJECUTIVO
        # Período de monitoreo
        total_dias = (df["datetime"].max() - df["datetime"].min()).days
        val_monitoreo = "Pasa" if total_dias >= 7 else "No pasa"

        # Frecuencia
        frec_min = round(df["FA"].min(), 2) if "FA" in df.columns else 0
        frec_max = round(df["FA"].max(), 2) if "FA" in df.columns else 0
        if "FA" in df.columns:
            val_frecuencia = "Pasa" if (frec_min >= 59.00 and frec_max <= 61.00) else "No pasa"
        else:
            val_frecuencia = "NA"

        # --- LIMITES DINÁMICOS EN BASE A CONFIG.PY (Tolerancia ±5% Promedios) ---
        limite_prom_inferior = round(config.VOLTAJE_NOMINAL_FASE * 0.95, 2)
        limite_prom_superior = round(config.VOLTAJE_NOMINAL_FASE * 1.05, 2)

        v1_prom = round(df["UA"].mean(), 2) if "UA" in df.columns else 0
        v2_prom = round(df["UB"].mean(), 2) if "UB" in df.columns else 0
        v3_prom = round(df["UC"].mean(), 2) if "UC" in df.columns else 0

        val_v1 = "Pasa" if limite_prom_inferior <= v1_prom <= limite_prom_superior else "No pasa"
        val_v2 = "Pasa" if limite_prom_inferior <= v2_prom <= limite_prom_superior else "No pasa"
        val_v3 = "Pasa" if limite_prom_inferior <= v3_prom <= limite_prom_superior else "No pasa"

        # Desbalances max o promedio (P95)
        desb_v = round(df["U-Unb-Neg"].quantile(0.95), 3) if "U-Unb-Neg" in df.columns else 0
        val_desb_v = "Pasa" if desb_v <= 2.0 else "No pasa"

        desb_i = round(df["I-Unb-Neg"].quantile(0.95), 2) if "I-Unb-Neg" in df.columns else 0
        val_desb_i = "Pasa" if desb_i <= 10.0 else "No pasa"

        # Máximos y mínimos absolutos registrados
        max_v = round(df[["UA", "UB", "UC"]].max().max(), 2) if "UA" in df.columns else 0
        min_v = round(df[["UA", "UB", "UC"]].min().min(), 2) if "UA" in df.columns else 0

        # Tolerancias absolutas del ±10% para eventos instantáneos
        limite_sobretension = round(config.VOLTAJE_NOMINAL_FASE * 1.10, 2)
        limite_subtension = round(config.VOLTAJE_NOMINAL_FASE * 0.90, 2)

        val_max_v = "Pasa" if max_v <= limite_sobretension else "No pasa"
        val_min_v = "Pasa" if min_v >= limite_subtension else "No pasa"

        # Conteo de eventos dinámico (Optimizado como Series independientes para evitar PerformanceWarning)
        eventos_sobretension = 0
        eventos_subtension = 0

        if "UA" in df.columns and "UB" in df.columns and "UC" in df.columns:
            serie_sobretension = (df["UA"] > limite_sobretension) | (df["UB"] > limite_sobretension) | (
                        df["UC"] > limite_sobretension)
            serie_subtension = (df["UA"] < limite_subtension) | (df["UB"] < limite_subtension) | (
                        df["UC"] < limite_subtension)

            # Multiplicar por 1 reemplaza a .astype(int) de forma nativa
            eventos_sobretension = int(((serie_sobretension * 1).diff() == 1).sum())
            eventos_subtension = int(((serie_subtension * 1).diff() == 1).sum())

            if serie_sobretension.iloc[0]: eventos_sobretension += 1
            if serie_subtension.iloc[0]: eventos_subtension += 1

        # Corrientes máximas
        i1_max = round(df["IA"].max(), 2) if "IA" in df.columns else 0
        i2_max = round(df["IB"].max(), 2) if "IB" in df.columns else 0
        i3_max = round(df["IC"].max(), 2) if "IC" in df.columns else 0

        # Factores de Cresta y Factor K promedios
        cf1_prom = round(df["IA-CF"].mean(), 2) if "IA-CF" in df.columns else 0
        cf2_prom = round(df["IB-CF"].mean(), 2) if "IB-CF" in df.columns else 0
        cf3_prom = round(df["IC-CF"].mean(), 2) if "IC-CF" in df.columns else 0

        kf1_prom = round(df["IA-KF"].mean(), 2) if "IA-KF" in df.columns else 0
        kf2_prom = round(df["IB-KF"].mean(), 2) if "IB-KF" in df.columns else 0
        kf3_prom = round(df["IC-KF"].mean(), 2) if "IC-KF" in df.columns else 0

        # Demandas y potencias máximas
        demanda_max = round(df["PSum"].max(), 2) if "PSum" in df.columns else 0
        aparente_max = round(df["SSum"].max(), 2) if "SSum" in df.columns else 0
        reactiva_max = round(df["QSum"].max(), 2) if "QSum" in df.columns else 0

        # Factor de potencia promedio
        fp_prom = round(df["PFTotal"].mean(), 2) if "PFTotal" in df.columns else 0
        if "PFTotal" in df.columns:
            pct_tiempo_fp = (df["PFTotal"] >= 0.97).sum() / len(df) * 100
            val_fp = "Pasa" if pct_tiempo_fp >= 95 else "No pasa"
        else:
            val_fp = "NA"

        # Construcción de tu matriz_resumen original
        matriz_resumen = [
            ["Período de monitoreo", f"{total_dias} días", "≥ 7 días (recomendado)", val_monitoreo],
            ["Tensión nominal del sistema",
             f"{int(config.VOLTAJE_NOMINAL_LINEA)}/{int(config.VOLTAJE_NOMINAL_FASE)} VAC",
             "Transformador Principal (Y)", "Informativo"],
            ["Frecuencia del sistema", f"Mín: {frec_min} Hz / Máx: {frec_max} Hz", "59.00 - 61.00 Hz", val_frecuencia],
            ["Voltaje promedio Fase 1-N", f"{v1_prom} V", f"±5% ({limite_prom_inferior} V - {limite_prom_superior} V)",
             val_v1],
            ["Voltaje promedio Fase 2-N", f"{v2_prom} V", f"±5% ({limite_prom_inferior} V - {limite_prom_superior} V)",
             val_v2],
            ["Voltaje promedio Fase 3-N", f"{v3_prom} V", f"±5% ({limite_prom_inferior} V - {limite_prom_superior} V)",
             val_v3],
            ["Desbalance de voltaje", f"{desb_v} %", "≤ 2 %", val_desb_v],
            ["Máxima sobretensión registrada", f"{max_v} V", f"≤ 110 % Vn ({limite_sobretension} V)", val_max_v],
            ["Mínima subtensión registrada", f"{min_v} V", f"≥ 90 % Vn ({limite_subtension} V)", val_min_v],
            ["Corriente máxima Fase 1", f"{i1_max} A", "Según capacidad del trafo", "NA"],
            ["Corriente máxima Fase 2", f"{i2_max} A", "Según capacidad del trafo", "NA"],
            ["Corriente máxima Fase 3", f"{i3_max} A", "Según capacidad del trafo", "NA"],
            ["Desbalance de corriente", f"{desb_i} %", "≤ 10 % (recomendado)", val_desb_i],
            ["Factor de Cresta promedio Fase 1", f"{cf1_prom}", "1.41 (Onda Senoidal Pura)", "Informativo"],
            ["Factor de Cresta promedio Fase 2", f"{cf2_prom}", "1.41 (Onda Senoidal Pura)", "Informativo"],
            ["Factor de Cresta promedio Fase 3", f"{cf3_prom}", "1.41 (Onda Senoidal Pura)", "Informativo"],
            ["Factor K promedio Fase 1", f"{kf1_prom}", "1.00 (Carga Lineal Ideal)", "Informativo"],
            ["Factor K promedio Fase 2", f"{kf2_prom}", "1.00 (Carga Lineal Ideal)", "Informativo"],
            ["Factor K promedio Fase 3", f"{kf3_prom}", "1.00 (Carga Lineal Ideal)", "Informativo"],
            ["Demanda máxima", f"{demanda_max} kW", "Informativo", "NA"],
            ["Potencia aparente máxima", f"{aparente_max} kVA", "Según capacidad del trafo", "NA"],
            ["Factor de potencia promedio", f"{fp_prom}", "≥ 0.97 (95% del tiempo)", val_fp],
            ["Potencia reactiva máxima", f"{reactiva_max} kVAR", "Informativo", "NA"],
            ["Eventos de sobretensión", f"{eventos_sobretension}", "Informativo", "NA"],
            ["Eventos de subtensión", f"{eventos_subtension}", "Informativo", "NA"]
        ]

        # 6. ENSAMBLAR EL ARCHIVO WORD USANDO REPORT_UTILS
        doc = Document()
        doc.add_heading('Reporte Técnico de Calidad de Energía', level=0)

        doc.add_paragraph("Este informe presenta un análisis exhaustivo del comportamiento del sistema "
                          "eléctrico ordenado por subsistemas, evaluado frente a requerimientos normativos.")

        # --- SECCIÓN 1: RESUMEN EJECUTIVO ---
        doc.add_heading('1. Resumen Ejecutivo de Requerimientos', level=1)
        doc.add_paragraph("La siguiente tabla condensa los hallazgos principales del periodo de medición, "
                          "asociándolos de manera directa a los criterios de validación reglamentarios:")
        # Llamamos a tu función personalizada del módulo report_utils
        report_utils.insertar_tabla_resumen_ejecutivo(doc, matriz_resumen)

        # --- SECCIÓN 2: ANÁLISIS POR BLOQUES Y GRÁFICAS ---
        doc.add_heading('2. Análisis de Perfiles y Comportamiento Temporal', level=1)

        for bloque in config.configuracion_reporte:
            titulo_bloque = bloque["titulo"]

            # Si la gráfica se generó exitosamente en memoria para este bloque
            if titulo_bloque in imagenes_memoria:
                doc.add_heading(titulo_bloque, level=2)

                # A) Insertar los Requerimientos Normativos del bloque si existen en tu config.py
                if "requerimientos" in bloque:
                    report_utils.insertar_tabla_requerimientos_horizontal(doc, bloque["requerimientos"])

                # B) Insertar la Tabla de Estadísticas de las variables de este bloque
                variables_del_bloque = [col_csv for col_csv, _ in bloque["variables"]]
                report_utils.insertar_tabla_superior_personalizada(doc, estadisticas, variables_del_bloque)

                # C) Insertar la Imagen directamente desde el flujo de memoria RAM
                doc.add_paragraph("Comportamiento temporal registrado:")
                doc.add_picture(imagenes_memoria[titulo_bloque], width=Inches(6.0))
                doc.add_paragraph()  # Espacio estético

        # 7. GUARDAR EL WORD EN MEMORIA Y RETORNAR AL USUARIO
        buffer_word = io.BytesIO()
        doc.save(buffer_word)
        buffer_word.seek(0)

        # Limpiar el nombre original reemplazando el .csv por .docx
        nombre_limpio = archivo.filename.rsplit('.', 1)[0]
        nombre_salida = f"Reporte_Calidad_{nombre_limpio}.docx"

        headers = {
            'Content-Disposition': f'attachment; filename="{nombre_salida}"'
        }

        return StreamingResponse(
            buffer_word,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=headers
        )

    except Exception as e:
        # Cualquier error interno del script será reportado limpiamente al frontend
        raise HTTPException(status_code=500, detail=f"Error procesando el análisis: {str(e)}")
