
from fastapi import FastAPI, HTTPException, Body
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware

import pandas as pd
import matplotlib

# Obliga a matplotlib a no usar interfaces gráficas de escritorio
# (crucial para servidores)
matplotlib.use("Agg")

import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches

import io

# Importaciones de tus módulos locales
import config
import report_utils


app = FastAPI(title="Procesador Avanzado de Calidad de Energía")


# ============================================================
# CORS
# ============================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# NORMALIZACIÓN DEL JSON
# ============================================================

def normalizar_datos(datos):

    df = pd.DataFrame(datos)

    if df.empty:
        raise HTTPException(
            status_code=400,
            detail="El JSON recibido no contiene registros."
        )

    mapeo = {
        "vl1": "UA",
        "vl2": "UB",
        "vl3": "UC",

        "vl12": "UAB",
        "vl23": "UBC",
        "vl13": "UCA",

        "hz": "FA",

        "unbV_neg": "U-Unb-Neg",
        "unbI_neg": "I-Unb-Neg",

        "thd_v1": "UA-THD",
        "thd_v2": "UB-THD",
        "thd_v3": "UC-THD",

        "thd_i1": "IA-THD",
        "thd_i2": "IB-THD",
        "thd_i3": "IC-THD",

        "i1": "IA",
        "i2": "IB",
        "i3": "IC",

        "fpt": "PFTotal",

        "kwt": "PSum",
        "kvart": "QSum",
        "kvat": "SSum",

        "K_L1": "IA-KF",
        "K_L2": "IB-KF",
        "K_L3": "IC-KF",
    }

    for origen, destino in mapeo.items():
        if origen in df.columns:
            df[destino] = df[origen]

    df["datetime"] = pd.to_datetime(
        df["listeningdate"],
        errors="coerce"
    )

    df = df.dropna(subset=["datetime"])

    df = df.sort_values("datetime").reset_index(drop=True)

    return df


# ============================================================
# ENDPOINT
# ============================================================

@app.post("/pqanalisys")
@app.post("/pqanalisys/")
async def analizar_historico(
    datos: list[dict] = Body(...)
):

    try:

        # ====================================================
        # 1. VALIDAR JSON
        # ====================================================

        if not datos:
            raise HTTPException(
                status_code=400,
                detail="El JSON recibido está vacío."
            )

        # ====================================================
        # 2. CONVERTIR JSON A DATAFRAME
        # ====================================================

        df = normalizar_datos(datos)

        # ====================================================
        # 3. COLUMNAS VÁLIDAS SEGÚN CONFIG.PY
        # ====================================================

        columnas_validas = [
            c
            for c in config.columnas_seleccionadas
            if c in df.columns
        ]

        if not columnas_validas:

            raise HTTPException(
                status_code=400,
                detail=(
                    "No se encontraron las columnas seleccionadas "
                    "en el JSON recibido."
                )
            )

        # ====================================================
        # 4. CONVERTIR COLUMNAS NUMÉRICAS
        # ====================================================

        df[columnas_validas] = df[columnas_validas].apply(
            pd.to_numeric,
            errors="coerce"
        )

        # ====================================================
        # 5. TABLA GENERAL DE ESTADÍSTICAS
        # ====================================================

        estadisticas = pd.DataFrame({
            "Mínimo": df[columnas_validas].min(),
            "Máximo": df[columnas_validas].max(),
            "Promedio": df[columnas_validas].mean(),
            "Desv. estándar": df[columnas_validas].std(),
            "P95": df[columnas_validas].quantile(0.95),
            "P99": df[columnas_validas].quantile(0.99)
        }).round(2)

        # ====================================================
        # 6. GENERACIÓN DE GRÁFICAS
        # ====================================================

        imagenes_memoria = {}

        for bloque in config.configuracion_reporte:

            plt.figure(figsize=(12, 5))

            grafico_algo = False

            for columna, etiqueta in bloque["variables"]:

                if columna in df.columns:

                    plt.plot(
                        df["datetime"],
                        df[columna],
                        label=etiqueta
                    )

                    grafico_algo = True

            if grafico_algo:

                plt.title(bloque["titulo"])
                plt.xlabel("Fecha y hora")
                plt.ylabel(bloque["ylabel"])

                plt.legend()
                plt.grid(True)

                plt.xticks(rotation=45)

                plt.tight_layout()

                # Guardar gráfica en memoria
                img_buf = io.BytesIO()

                plt.savefig(
                    img_buf,
                    format="png",
                    dpi=250,
                    bbox_inches="tight"
                )

                img_buf.seek(0)

                imagenes_memoria[
                    bloque["titulo"]
                ] = img_buf

            plt.close()

        # ====================================================
        # 7. RESUMEN EJECUTIVO
        # ====================================================

        total_dias = (
            df["datetime"].max() -
            df["datetime"].min()
        ).days

        val_monitoreo = (
            "Pasa"
            if total_dias >= 7
            else "No pasa"
        )

        # ----------------------------------------------------
        # Frecuencia
        # ----------------------------------------------------

        frec_min = (
            round(df["FA"].min(), 2)
            if "FA" in df.columns
            else 0
        )

        frec_max = (
            round(df["FA"].max(), 2)
            if "FA" in df.columns
            else 0
        )

        if "FA" in df.columns:

            val_frecuencia = (
                "Pasa"
                if (
                    frec_min >= 59.00
                    and frec_max <= 61.00
                )
                else "No pasa"
            )

        else:

            val_frecuencia = "NA"

        # ----------------------------------------------------
        # Límites de voltaje
        # ----------------------------------------------------

        limite_prom_inferior = round(
            config.VOLTAJE_NOMINAL_FASE * 0.95,
            2
        )

        limite_prom_superior = round(
            config.VOLTAJE_NOMINAL_FASE * 1.05,
            2
        )

        v1_prom = (
            round(df["UA"].mean(), 2)
            if "UA" in df.columns
            else 0
        )

        v2_prom = (
            round(df["UB"].mean(), 2)
            if "UB" in df.columns
            else 0
        )

        v3_prom = (
            round(df["UC"].mean(), 2)
            if "UC" in df.columns
            else 0
        )

        val_v1 = (
            "Pasa"
            if limite_prom_inferior <= v1_prom <= limite_prom_superior
            else "No pasa"
        )

        val_v2 = (
            "Pasa"
            if limite_prom_inferior <= v2_prom <= limite_prom_superior
            else "No pasa"
        )

        val_v3 = (
            "Pasa"
            if limite_prom_inferior <= v3_prom <= limite_prom_superior
            else "No pasa"
        )

        # ----------------------------------------------------
        # Desbalance de voltaje
        # ----------------------------------------------------

        desb_v = (
            round(
                df["U-Unb-Neg"].quantile(0.95),
                3
            )
            if "U-Unb-Neg" in df.columns
            else 0
        )

        val_desb_v = (
            "Pasa"
            if desb_v <= 2.0
            else "No pasa"
        )

        # ----------------------------------------------------
        # Desbalance de corriente
        # ----------------------------------------------------

        desb_i = (
            round(
                df["I-Unb-Neg"].quantile(0.95),
                2
            )
            if "I-Unb-Neg" in df.columns
            else 0
        )

        val_desb_i = (
            "Pasa"
            if desb_i <= 10.0
            else "No pasa"
        )

        # ----------------------------------------------------
        # Máximos y mínimos de voltaje
        # ----------------------------------------------------

        columnas_voltaje = [
            c
            for c in ["UA", "UB", "UC"]
            if c in df.columns
        ]

        if columnas_voltaje:

            max_v = round(
                df[columnas_voltaje].max().max(),
                2
            )

            min_v = round(
                df[columnas_voltaje].min().min(),
                2
            )

        else:

            max_v = 0
            min_v = 0

        # ----------------------------------------------------
        # Límites absolutos
        # ----------------------------------------------------

        limite_sobretension = round(
            config.VOLTAJE_NOMINAL_FASE * 1.10,
            2
        )

        limite_subtension = round(
            config.VOLTAJE_NOMINAL_FASE * 0.90,
            2
        )

        val_max_v = (
            "Pasa"
            if max_v <= limite_sobretension
            else "No pasa"
        )

        val_min_v = (
            "Pasa"
            if min_v >= limite_subtension
            else "No pasa"
        )

        # ====================================================
        # 8. EVENTOS DE TENSIÓN
        # ====================================================

        eventos_sobretension = 0
        eventos_subtension = 0

        if all(
            c in df.columns
            for c in ["UA", "UB", "UC"]
        ):

            serie_sobretension = (
                (df["UA"] > limite_sobretension)
                |
                (df["UB"] > limite_sobretension)
                |
                (df["UC"] > limite_sobretension)
            )

            serie_subtension = (
                (df["UA"] < limite_subtension)
                |
                (df["UB"] < limite_subtension)
                |
                (df["UC"] < limite_subtension)
            )

            eventos_sobretension = int(
                ((serie_sobretension * 1).diff() == 1).sum()
            )

            eventos_subtension = int(
                ((serie_subtension * 1).diff() == 1).sum()
            )

            if serie_sobretension.iloc[0]:
                eventos_sobretension += 1

            if serie_subtension.iloc[0]:
                eventos_subtension += 1

        # ====================================================
        # 9. CORRIENTES MÁXIMAS
        # ====================================================

        i1_max = (
            round(df["IA"].max(), 2)
            if "IA" in df.columns
            else 0
        )

        i2_max = (
            round(df["IB"].max(), 2)
            if "IB" in df.columns
            else 0
        )

        i3_max = (
            round(df["IC"].max(), 2)
            if "IC" in df.columns
            else 0
        )

        # ====================================================
        # 10. FACTOR DE CRESTA
        # ====================================================

        cf1_prom = (
            round(df["IA-CF"].mean(), 2)
            if "IA-CF" in df.columns
            else 0
        )

        cf2_prom = (
            round(df["IB-CF"].mean(), 2)
            if "IB-CF" in df.columns
            else 0
        )

        cf3_prom = (
            round(df["IC-CF"].mean(), 2)
            if "IC-CF" in df.columns
            else 0
        )

        # ====================================================
        # 11. FACTOR K
        # ====================================================

        kf1_prom = (
            round(df["IA-KF"].mean(), 2)
            if "IA-KF" in df.columns
            else 0
        )

        kf2_prom = (
            round(df["IB-KF"].mean(), 2)
            if "IB-KF" in df.columns
            else 0
        )

        kf3_prom = (
            round(df["IC-KF"].mean(), 2)
            if "IC-KF" in df.columns
            else 0
        )

        # ====================================================
        # 12. DEMANDAS Y POTENCIAS
        # ====================================================

        demanda_max = (
            round(df["PSum"].max(), 2)
            if "PSum" in df.columns
            else 0
        )

        aparente_max = (
            round(df["SSum"].max(), 2)
            if "SSum" in df.columns
            else 0
        )

        reactiva_max = (
            round(df["QSum"].max(), 2)
            if "QSum" in df.columns
            else 0
        )

        # ====================================================
        # 13. FACTOR DE POTENCIA
        # ====================================================

        fp_prom = (
            round(df["PFTotal"].mean(), 2)
            if "PFTotal" in df.columns
            else 0
        )

        if "PFTotal" in df.columns:

            pct_tiempo_fp = (
                (df["PFTotal"] >= 0.97).sum()
                / len(df)
                * 100
            )

            val_fp = (
                "Pasa"
                if pct_tiempo_fp >= 95
                else "No pasa"
            )

        else:

            val_fp = "NA"

        # ====================================================
        # 14. MATRIZ DEL RESUMEN
        # ====================================================

        matriz_resumen = [

            [
                "Período de monitoreo",
                f"{total_dias} días",
                "≥ 7 días (recomendado)",
                val_monitoreo
            ],

            [
                "Tensión nominal del sistema",
                (
                    f"{int(config.VOLTAJE_NOMINAL_LINEA)}/"
                    f"{int(config.VOLTAJE_NOMINAL_FASE)} VAC"
                ),
                "Transformador Principal (Y)",
                "Informativo"
            ],

            [
                "Frecuencia del sistema",
                f"Mín: {frec_min} Hz / Máx: {frec_max} Hz",
                "59.00 - 61.00 Hz",
                val_frecuencia
            ],

            [
                "Voltaje promedio Fase 1-N",
                f"{v1_prom} V",
                (
                    f"±5% "
                    f"({limite_prom_inferior} V - "
                    f"{limite_prom_superior} V)"
                ),
                val_v1
            ],

            [
                "Voltaje promedio Fase 2-N",
                f"{v2_prom} V",
                (
                    f"±5% "
                    f"({limite_prom_inferior} V - "
                    f"{limite_prom_superior} V)"
                ),
                val_v2
            ],

            [
                "Voltaje promedio Fase 3-N",
                f"{v3_prom} V",
                (
                    f"±5% "
                    f"({limite_prom_inferior} V - "
                    f"{limite_prom_superior} V)"
                ),
                val_v3
            ],

            [
                "Desbalance de voltaje",
                f"{desb_v} %",
                "≤ 2 %",
                val_desb_v
            ],

            [
                "Máxima sobretensión registrada",
                f"{max_v} V",
                (
                    f"≤ 110 % Vn "
                    f"({limite_sobretension} V)"
                ),
                val_max_v
            ],

            [
                "Mínima subtensión registrada",
                f"{min_v} V",
                (
                    f"≥ 90 % Vn "
                    f"({limite_subtension} V)"
                ),
                val_min_v
            ],

            [
                "Corriente máxima Fase 1",
                f"{i1_max} A",
                "Según capacidad del trafo",
                "NA"
            ],

            [
                "Corriente máxima Fase 2",
                f"{i2_max} A",
                "Según capacidad del trafo",
                "NA"
            ],

            [
                "Corriente máxima Fase 3",
                f"{i3_max} A",
                "Según capacidad del trafo",
                "NA"
            ],

            [
                "Desbalance de corriente",
                f"{desb_i} %",
                "≤ 10 % (recomendado)",
                val_desb_i
            ],

            [
                "Factor de Cresta promedio Fase 1",
                f"{cf1_prom}",
                "1.41 (Onda Senoidal Pura)",
                "Informativo"
            ],

            [
                "Factor de Cresta promedio Fase 2",
                f"{cf2_prom}",
                "1.41 (Onda Senoidal Pura)",
                "Informativo"
            ],

            [
                "Factor de Cresta promedio Fase 3",
                f"{cf3_prom}",
                "1.41 (Onda Senoidal Pura)",
                "Informativo"
            ],

            [
                "Factor K promedio Fase 1",
                f"{kf1_prom}",
                "1.00 (Carga Lineal Ideal)",
                "Informativo"
            ],

            [
                "Factor K promedio Fase 2",
                f"{kf2_prom}",
                "1.00 (Carga Lineal Ideal)",
                "Informativo"
            ],

            [
                "Factor K promedio Fase 3",
                f"{kf3_prom}",
                "1.00 (Carga Lineal Ideal)",
                "Informativo"
            ],

            [
                "Demanda máxima",
                f"{demanda_max} kW",
                "Informativo",
                "NA"
            ],

            [
                "Potencia aparente máxima",
                f"{aparente_max} kVA",
                "Según capacidad del trafo",
                "NA"
            ],

            [
                "Factor de potencia promedio",
                f"{fp_prom}",
                "≥ 0.97 (95% del tiempo)",
                val_fp
            ],

            [
                "Potencia reactiva máxima",
                f"{reactiva_max} kVAR",
                "Informativo",
                "NA"
            ],

            [
                "Eventos de sobretensión",
                f"{eventos_sobretension}",
                "Informativo",
                "NA"
            ],

            [
                "Eventos de subtensión",
                f"{eventos_subtension}",
                "Informativo",
                "NA"
            ]
        ]

        # ====================================================
        # 15. CREAR DOCUMENTO WORD
        # ====================================================

        doc = Document()

        doc.add_heading(
            "Reporte Técnico de Calidad de Energía",
            level=0
        )

        doc.add_paragraph(
            "Este informe presenta un análisis exhaustivo "
            "del comportamiento del sistema eléctrico "
            "ordenado por subsistemas, evaluado frente "
            "a requerimientos normativos."
        )

        # ====================================================
        # SECCIÓN 1
        # ====================================================

        doc.add_heading(
            "1. Resumen Ejecutivo de Requerimientos",
            level=1
        )

        doc.add_paragraph(
            "La siguiente tabla condensa los hallazgos "
            "principales del periodo de medición, "
            "asociándolos de manera directa a los "
            "criterios de validación reglamentarios:"
        )

        report_utils.insertar_tabla_resumen_ejecutivo(
            doc,
            matriz_resumen
        )

        # ====================================================
        # SECCIÓN 2
        # ====================================================

        doc.add_heading(
            "2. Análisis de Perfiles y Comportamiento Temporal",
            level=1
        )

        for bloque in config.configuracion_reporte:

            titulo_bloque = bloque["titulo"]

            if titulo_bloque not in imagenes_memoria:
                continue

            doc.add_heading(
                titulo_bloque,
                level=2
            )

            # Requerimientos normativos
            if "requerimientos" in bloque:

                report_utils.insertar_tabla_requerimientos_horizontal(
                    doc,
                    bloque["requerimientos"]
                )

            # Estadísticas
            variables_del_bloque = [
                columna
                for columna, _ in bloque["variables"]
            ]

            report_utils.insertar_tabla_superior_personalizada(
                doc,
                estadisticas,
                variables_del_bloque
            )

            # Gráfica
            doc.add_paragraph(
                "Comportamiento temporal registrado:"
            )

            doc.add_picture(
                imagenes_memoria[titulo_bloque],
                width=Inches(6.0)
            )

            doc.add_paragraph()

        # ====================================================
        # 16. GUARDAR WORD EN MEMORIA
        # ====================================================

        buffer_word = io.BytesIO()

        doc.save(buffer_word)

        buffer_word.seek(0)

        # ====================================================
        # 17. NOMBRE DEL ARCHIVO
        # ====================================================

        fecha_inicio = (
            df["datetime"].min()
            .strftime("%Y%m%d")
        )

        fecha_fin = (
            df["datetime"].max()
            .strftime("%Y%m%d")
        )

        nombre_salida = (
            f"Reporte_Calidad_"
            f"{fecha_inicio}_"
            f"{fecha_fin}.docx"
        )

        # ====================================================
        # 18. RESPUESTA HTTP
        # ====================================================

        headers = {
            "Content-Disposition":
                f'attachment; filename="{nombre_salida}"'
        }

        return StreamingResponse(
            buffer_word,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
            headers=headers
        )

    except HTTPException:
        raise

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=(
                "Error procesando el análisis: "
                f"{str(e)}"
            )
        )