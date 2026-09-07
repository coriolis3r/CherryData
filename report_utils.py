# report_utils.py

from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

HEX_ENCABEZADO = "162F4D"
HEX_GRIS_CLARO = "F2F2F2"
COLOR_TEXTO = RGBColor(22, 47, 77)


def colorear_celda(celda, hex_color):
    """Inserta propiedades XML para sombrear el fondo de una celda."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    celda._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def insertar_tabla_requerimientos_horizontal(doc_obj, dict_req):
    """Crea una tabla horizontal con los encabezados normativos en la parte superior."""
    tabla = doc_obj.add_table(rows=2, cols=3)
    tabla.style = 'Table Grid'

    # Encabezados (Fila 0)
    hdr_cells = tabla.rows[0].cells
    for index, encabezado_titulo in enumerate(dict_req.keys()):
        celda = hdr_cells[index]
        colorear_celda(celda, HEX_ENCABEZADO)
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(encabezado_titulo))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Valores (Fila 1)
    val_cells = tabla.rows[1].cells
    for index, valor_requerimiento in enumerate(dict_req.values()):
        celda = val_cells[index]
        colorear_celda(celda, HEX_GRIS_CLARO)
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(valor_requerimiento))
        run.font.color.rgb = COLOR_TEXTO

    doc_obj.add_paragraph()


def insertar_tabla_superior_personalizada(doc_obj, df_estadisticas, lista_variables):
    """Inserta la tabla horizontal original con los cálculos del CSV."""
    # Filtrar únicamente las variables correspondientes a este bloque de la gráfica
    df_filtrado = df_estadisticas.loc[df_estadisticas.index.isin(lista_variables)]

    if df_filtrado.empty:
        return

    # Extraer el número entero de filas y columnas
    num_filas = df_filtrado.shape[0] + 1
    num_columnas = df_filtrado.shape[1] + 1

    tabla = doc_obj.add_table(rows=num_filas, cols=num_columnas)
    tabla.style = 'Table Grid'

    # --- ENCABEZADOS (Corregido: acceder a la fila 0) ---
    cabeceras = ["Variable"] + list(df_filtrado.columns)
    hdr_cells = tabla.rows[0].cells  # <--- Cambio clave aquí
    for index, texto in enumerate(cabeceras):
        celda = hdr_cells[index]
        colorear_celda(celda, HEX_ENCABEZADO)
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(texto))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # --- FILAS DE DATOS (Efecto Zebra) ---
    for idx_fila, (variable_nombre, fila_valores) in enumerate(df_filtrado.iterrows()):
        celdas = tabla.rows[idx_fila + 1].cells
        es_fila_gris = (idx_fila % 2 != 0)
        color_fondo = HEX_GRIS_CLARO if es_fila_gris else "FFFFFF"

        # Primera columna: Nombre de la Variable
        colorear_celda(celdas[0], color_fondo)
        p_var = celdas[0].paragraphs[0]
        run_var = p_var.add_run(str(variable_nombre))
        run_var.font.color.rgb = COLOR_TEXTO
        run_var.bold = True

        # Columnas numéricas
        for idx_col, valor in enumerate(fila_valores):
            celda_valor = celdas[idx_col + 1]
            colorear_celda(celda_valor, color_fondo)
            p_val = celda_valor.paragraphs[0]
            p_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run_val = p_val.add_run(str(valor))
            run_val.font.color.rgb = COLOR_TEXTO

    doc_obj.add_paragraph()


def insertar_tabla_resumen_ejecutivo(doc_obj, datos_resumen):
    """Crea una tabla horizontal de resumen ejecutivo con columnas fijas."""
    # Matriz: 1 fila de cabecera + filas de datos. 4 columnas fijas.
    num_filas = len(datos_resumen) + 1
    tabla = doc_obj.add_table(rows=num_filas, cols=4)
    tabla.style = 'Table Grid'

    # --- CABECERAS ---
    cabeceras = ["Parámetro", "Resultado", "Valor de ref", "Validación"]
    hdr_cells = tabla.rows[0].cells
    for index, texto in enumerate(cabeceras):
        celda = hdr_cells[index]
        colorear_celda(celda, HEX_ENCABEZADO)
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(texto))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # --- FILAS DE DATOS ---
    for idx_fila, fila in enumerate(datos_resumen):
        celdas = tabla.rows[idx_fila + 1].cells
        es_fila_gris = (idx_fila % 2 != 0)
        color_fondo = HEX_GRIS_CLARO if es_fila_gris else "FFFFFF"

        for idx_col, valor in enumerate(fila):
            celda = celdas[idx_col]
            colorear_celda(celda, color_fondo)
            p = celda.paragraphs[0]

            # Alineaciones estéticas por columna (Corregido)
            if idx_col == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif idx_col in [1, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = p.add_run(str(valor))
            run.font.color.rgb = COLOR_TEXTO

            # Destacar fallas en negrita si no pasa
            if idx_col == 3 and str(valor) == "No pasa":
                run.bold = True

    doc_obj.add_paragraph()

